#!/usr/bin/env python3
"""Generate the Cisco CCNP ENCOR Lesson Plan (LP) DOCX in the Tertiary house format.

Cover page + Document Version Control Record + auto TOC + Arial 11pt body +
colour-coded 5-day schedule tables (9:30am-6:30pm, 8 training hours/day, 1h
lunch, tea within, final assessment Day 5 from 4:30pm).

Slide numbers are read DIRECTLY out of the generated deck, so the LP can never
drift from the PPT: rebuild the deck, rebuild the LP, and the references follow.
"""
import os, sys, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5; from data_domain6 import DOMAIN6
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5+DOMAIN6
import prodoc

def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
HEADER_FILL="1F6FEB"; TOPIC_FILL="E8F0FE"; BREAK_FILL="FFF4E5"; LUNCH_FILL="FDE9D9"; ASSESS_FILL="E8F7EE"

# ---------------------------------------------- read slide numbers out of the deck
def read_slide_map():
    """Return (sections, labs, total) read from the CURRENT deck, so LP == PPT."""
    from pptx import Presentation
    deck=os.path.join(REPO,"courseware",f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
    if not os.path.exists(deck):
        raise SystemExit(f"Deck not found: {deck}\nBuild the PPT before the LP so slide numbers align.")
    p=Presentation(deck); sec={}; lab={}
    for i,s in enumerate(p.slides):
        t=[sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        j=" || ".join(t)
        m=re.match(r"TOPIC (\d\d) \|\|",j)
        if m and "SCOPE" not in j: sec["T"+m.group(1)]=i+1
        m2=re.search(r"\|\| LAB (\d+) \|\|",j)
        if m2 and "HANDS-ON" in j: lab[int(m2.group(1))]=i+1
        if "COURSE ADMINISTRATION ||" in j: sec["ADMIN"]=i+1
        if "CORE CONCEPTS ||" in j: sec["CORE"]=i+1
        if "WRAP-UP ||" in j: sec["WRAP"]=i+1
    return sec,lab,len(p.slides)

SEC,LAB,TOTAL_SLIDES = read_slide_map()

def slides_for(nums):
    """Slide range string covering the given lab numbers."""
    got=[LAB[n] for n in nums if n in LAB]
    if not got: return ""
    lo=min(got); hi_lab=max(nums)
    hi=LAB.get(hi_lab,lo)+2      # overview + process map + verify
    return f"Slides {lo}–{hi}"

def lab_titles(nums):
    return "; ".join(f"Lab {a['num']}: {a['title']}" for a in ACT if a['num'] in nums)

def block(nums, lead=""):
    """Activity text = optional lead-in + lab titles + the slide reference."""
    parts=[p for p in [lead, "Hands-on: "+lab_titles(nums) if nums else ""] if p]
    txt=" ".join(parts)
    sr=slides_for(nums)
    return f"{txt} ({sr})" if sr else txt

# ------------------------------------------------ schedule (single source of truth for timing)
# (start, end, minutes, kind, activity_text)  kind: admin/topic/lab/break/lunch/assess/recap
SCHEDULE = {
 1: (C.DAY_THEMES[1], [
    ("9:30","10:15",45,"admin",f"Welcome, trainer and learner introductions, ground rules, learning outcomes, course outline, Skills Framework alignment and mandatory digital attendance (AM) (Slides {SEC['ADMIN']}–{SEC['CORE']-1})"),
    ("10:15","11:15",60,"topic",f"Core concepts — the ENCOR 350-401 blueprint, the three planes, switching mechanisms, the forwarding tables and the campus design decision (Slides {SEC['CORE']}–{SEC['T01']-1})"),
    ("11:15","11:30",15,"break","Tea break"),
    ("11:30","13:00",90,"topic",f"Topic 1 — Enterprise Architecture: hierarchical design, high availability, wireless design principles, SD-WAN and SD-Access planes, QoS components and policy (Slides {SEC['T01']}–{LAB[1]-1})"),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:45",105,"lab",block([1,2,3],"Digital attendance (PM).")),
    ("15:45","16:00",15,"break","Tea break"),
    ("16:00","18:15",135,"lab",block([4,5])),
    ("18:15","18:30",15,"recap","Day 1 recap, Q&A and consolidation of Topic 1"),
 ]),
 2: (C.DAY_THEMES[2], [
    ("9:30","9:45",15,"recap","Day 1 recap and mandatory digital attendance (AM)"),
    ("9:45","11:15",90,"topic",f"Topic 2 — Virtualization Technologies: Type 1 and Type 2 hypervisors, virtual machines, virtual switching, and the concepts behind VRF, GRE, IPsec, LISP and VXLAN (Slides {SEC['T02']}–{LAB[6]-1})"),
    ("11:15","11:30",15,"break","Tea break"),
    ("11:30","13:00",90,"lab",block([6])),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:45",105,"lab",block([7,8],"Digital attendance (PM).")),
    ("15:45","16:00",15,"break","Tea break"),
    ("16:00","18:15",135,"lab",block([9],"Consolidation of the data path virtualization objectives.")),
    ("18:15","18:30",15,"recap","Day 2 recap, Q&A and consolidation of Topic 2"),
 ]),
 3: (C.DAY_THEMES[3], [
    ("9:30","9:45",15,"recap","Day 2 recap and mandatory digital attendance (AM)"),
    ("9:45","11:15",90,"topic",f"Topic 3 — Network Infrastructure (Layer 2): 802.1Q trunking and DTP, EtherChannel with PAgP and LACP, STP, RSTP, MST and the Spanning Tree protection features (Slides {SEC['T03']}–{LAB[10]-1})"),
    ("11:15","11:30",15,"break","Tea break"),
    ("11:30","13:00",90,"lab",block([10,11])),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:45",105,"lab",block([12,13,14],"Digital attendance (PM).")),
    ("15:45","16:00",15,"break","Tea break"),
    ("16:00","18:15",135,"lab",block([15,16,17,18],"Topic 3 (Layer 3): EIGRP versus OSPF path selection, multiarea OSPFv2 and OSPFv3, eBGP and policy-based routing.")),
    ("18:15","18:30",15,"recap","Day 3 recap, Q&A and consolidation of the Layer 2 and Layer 3 objectives"),
 ]),
 4: (C.DAY_THEMES[4], [
    ("9:30","9:45",15,"recap","Day 3 recap and mandatory digital attendance (AM)"),
    ("9:45","11:15",90,"lab",block([19,20],"Topic 3 (wireless): Layer 1 concepts, AP modes and antennas, CAPWAP discovery and join, Layer 2 and Layer 3 roaming, and wireless segmentation.")),
    ("11:15","11:30",15,"break","Tea break"),
    ("11:30","13:00",90,"lab",block([21],"Topic 3 (IP services): NTP and PTP, NAT and PAT, HSRP, VRRP and GLBP, and multicast with RPF, PIM and IGMP.")),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:45",105,"lab",block([22,23,24,25],f"Digital attendance (PM). Topic 4 — Network Assurance: diagnostics, Flexible NetFlow, SPAN/RSPAN/ERSPAN, IP SLA, DNA Center workflows and NETCONF/RESTCONF (Slides {SEC['T04']}–{LAB[22]-1}).")),
    ("15:45","16:00",15,"break","Tea break"),
    ("16:00","18:15",135,"lab",block([26,27,28,29,30],f"Topic 5 — Network Security: device access control, AAA, ACLs, CoPP, REST API security, wireless security and security design (Slides {SEC['T05']}–{LAB[26]-1}).")),
    ("18:15","18:30",15,"recap","Day 4 recap, Q&A and consolidation of Topics 4 and 5"),
 ]),
 5: (C.DAY_THEMES[5], [
    ("9:30","9:45",15,"recap","Day 4 recap and mandatory digital attendance (AM)"),
    ("9:45","11:15",90,"lab",block([31,32],f"Topic 6 — Network Automation: interpreting Python, constructing valid JSON, YANG data modelling (Slides {SEC['T06']}–{LAB[31]-1}).")),
    ("11:15","11:30",15,"break","Tea break"),
    ("11:30","13:00",90,"lab",block([33,34])),
    ("13:00","14:00",60,"lunch","Lunch break"),
    ("14:00","15:15",75,"lab",block([35],"Digital attendance (PM).")),
    ("15:15","16:00",45,"topic",f"Course revision across all six exam domains, exam-technique guidance and the practice exam at exams.tertiaryinfotech.com (Slides {SEC['WRAP']}–{TOTAL_SLIDES})"),
    ("16:00","16:15",15,"recap","Course feedback and mandatory TRAQOM survey"),
    ("16:15","16:30",15,"assess","Briefing for Assessment and Assessment digital attendance"),
    ("16:30","17:30",60,"assess","Written Assessment (WA) — Short-Answer Questions (SAQ), 1 hour, open book"),
    ("17:30","18:30",60,"assess","Practical Performance (PP) — hands-on Cisco configuration and verification tasks, 1 hour, open book"),
 ]),
}

# ------------------------------------------------ build document
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)

prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("5.0","1 April 2024","Previous release — 5-day CCNP ENCOR lesson plan built on the legacy master trainer slides.",C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "Full rebuild against the Cisco ENCOR 350-401 v1.1 exam blueprint. Six topics realigned to the six exam domains with their published weightings; 35 hands-on Packet Tracer labs replacing the legacy discovery exercises; schedule rebuilt to 9:30am–6:30pm with 8 instructional hours per day; slide references now generated directly from the deck.",C.TRAINER),
])
prodoc.add_toc(doc)

def H(text,level=1):
    return doc.add_heading(text,level=level)

H("Course Information",1)
info=[("Course Title",C.TITLE),("WSQ Course Reference",C.COURSE_CODE),
      ("Training Provider",C.ORG+"  ("+C.UEN.replace('UEN: ','UEN ')+")"),
      ("Duration",f"{C.DAYS} days · 8 training hours per day (40 hours)"),
      ("Daily Timing","9:30 am – 6:30 pm (1-hour lunch; tea breaks counted within training time)"),
      ("Mode","Instructor-led, with hands-on Cisco Packet Tracer configuration and troubleshooting labs"),
      ("Certification Prepared For",f"Cisco {C.EXAM_CODE} — {C.EXAM_NAME}"),
      ("Skills Framework",f"TSC: {C.TSC_TITLE} ({C.TSC_CODE})"),
      ("Trainer",C.TRAINER),
      ("Trainer Slides",f"{C.SHORT_TITLE}-{C.VERSION}.pptx ({TOTAL_SLIDES} slides)")]
t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
for k,v in info:
    c=t.add_row().cells; c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
    prodoc._shade_cell(c[0],TOPIC_FILL)
    c[1].text=""; c[1].paragraphs[0].add_run(v).font.size=Pt(10)

H("Learning Outcomes",1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for lo in C.LEARNING_OUTCOMES:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(lo).font.size=Pt(10.5)

H("Skills Framework Mapping",1)
doc.add_paragraph(f"This course is mapped to the Skills Framework TSC {C.TSC_TITLE} ({C.TSC_CODE}).")
sf=doc.add_table(rows=0,cols=2); sf.style="Table Grid"
hdr=sf.add_row().cells
for i,htext in enumerate(["Knowledge (K)","Ability (A)"]):
    hdr[i].text=""; r=hdr[i].paragraphs[0].add_run(htext); r.bold=True; r.font.size=Pt(10)
    r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); prodoc._shade_cell(hdr[i],HEADER_FILL)
for (kc,kt),(ac,at) in zip(C.TSC_KNOWLEDGE,C.TSC_ABILITIES):
    cells=sf.add_row().cells
    cells[0].text=""; cells[0].paragraphs[0].add_run(f"{kc}: {kt}").font.size=Pt(9.5)
    cells[1].text=""; cells[1].paragraphs[0].add_run(f"{ac}: {at}").font.size=Pt(9.5)

H("Instructional Methods",1)
for m in ["Lecture and trainer-led explanation of each exam domain's concepts, supported by diagrams, comparison tables and worked examples in the trainer slides.",
          "Didactic questioning throughout each topic to check understanding and surface misconceptions before the labs.",
          "Demonstration of each configuration on Cisco IOS before learners attempt it.",
          "Practical hands-on labs in Cisco Packet Tracer — 35 labs, one or more for every exam objective in the blueprint.",
          "Guided troubleshooting exercises in which learners diagnose a deliberately broken configuration rather than building from blank."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(m).font.size=Pt(10.5)

H("Assessment",1)
for a in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],
          "Format: Open Book — course slides, Learner Guide and approved materials only.",
          "The final assessment is conducted on Day 5 from 4:30 pm, following the Briefing for Assessment and the Assessment digital attendance.",
          "Learners are assessed as Competent (C) or Not Yet Competent (NYC) on each instrument. An appeal process is available.",
          C.ASSESSMENT["note"]]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10.5)

def set_cell(cell,text,bold=False,size=9.5,color=None,fill=None,align=None):
    cell.text=""; p=cell.paragraphs[0]
    if align: p.alignment=align
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name="Arial"
    if color: r.font.color.rgb=color
    if fill: prodoc._shade_cell(cell,fill)

KIND_FILL={"topic":TOPIC_FILL,"break":BREAK_FILL,"lunch":LUNCH_FILL,"assess":ASSESS_FILL,
           "admin":"F3F5F8","recap":"F3F5F8","lab":None}

H("Course Schedule",1)
doc.add_paragraph("Each training day totals 8 instructional hours. The 1-hour lunch break is not counted as instructional time; tea breaks are counted within it. Slide references are to the trainer slide deck listed in Course Information.")
for day,(theme,rows) in SCHEDULE.items():
    H(f"Day {day} — {theme}",2)
    tbl=doc.add_table(rows=0,cols=3); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=tbl.add_row().cells
    for i,htext in enumerate(["Time","Duration","Topic / Activity"]):
        set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
    training=0
    for start,end,mins,kind,text in rows:
        cells=tbl.add_row().cells; fill=KIND_FILL.get(kind)
        set_cell(cells[0],f"{start}–{end}",bold=(kind in ("topic","assess")),size=9.5,fill=fill)
        set_cell(cells[1],f"{mins} min",size=9.5,fill=fill)
        set_cell(cells[2],text,bold=(kind in ("topic","assess")),size=9.5,fill=fill)
        if kind!="lunch": training+=mins
    for row in tbl.rows:
        row.cells[0].width=Inches(1.15); row.cells[1].width=Inches(0.9); row.cells[2].width=Inches(4.75)
    p=doc.add_paragraph(); r=p.add_run(f"Total instructional time: {training} minutes ({training//60} hours).")
    r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY
    assert training==480, f"Day {day} training minutes = {training}, expected 480"

H("Lab Reference — Aligned to the ENCOR 350-401 Exam Domains",1)
tt=doc.add_table(rows=0,cols=4); tt.style="Table Grid"
hdr=tt.add_row().cells
for i,htext in enumerate(["Topic / Exam domain","Weighting","TSC","Labs"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    cells=tt.add_row().cells
    set_cell(cells[0],f"Topic {tp['code']}: {tp['title']}  ({tp['exam']})",bold=True,size=9.5,fill=TOPIC_FILL)
    set_cell(cells[1],tp["weighting"],size=9.5,fill=TOPIC_FILL)
    set_cell(cells[2],tp["tsc"],size=9.5,fill=TOPIC_FILL)
    set_cell(cells[3],f"Labs {acts[0]['num']}–{acts[-1]['num']}",size=9.5)

H("Lab Index with Slide References",1)
lt=doc.add_table(rows=0,cols=4); lt.style="Table Grid"
hdr=lt.add_row().cells
for i,htext in enumerate(["Lab","Title","Exam objective","Slide"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for a in ACT:
    cells=lt.add_row().cells
    set_cell(cells[0],str(a["num"]),bold=True,size=9)
    set_cell(cells[1],a["title"],size=9)
    set_cell(cells[2],a["objective"],size=8.5)
    set_cell(cells[3],str(LAB.get(a["num"],"—")),size=9)
for row in lt.rows:
    row.cells[0].width=Inches(0.45); row.cells[1].width=Inches(2.2)
    row.cells[2].width=Inches(3.55); row.cells[3].width=Inches(0.6)

H("Resources Required",1)
for r_ in ["Cisco Packet Tracer 8.2 or later installed on each learner workstation (free from the Cisco Networking Academy).",
           "The lab pack for this course — one folder per lab containing the topology file, the addressing table and the mock data set.",
           "Python 3.10 or later with the requests library, for the Topic 6 automation labs.",
           "Internet access for the REST API labs and the practice exam at exams.tertiaryinfotech.com.",
           "Projector or large display for trainer demonstrations, and a whiteboard for topology sketches."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(r_).font.size=Pt(10.5)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
doc.save(OUT)
print("Saved",OUT,f"(slide map read from deck: {TOTAL_SLIDES} slides)")
